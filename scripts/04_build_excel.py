import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
CLEAN = BASE_DIR / "data/clean"
OUT = BASE_DIR / "excel/Swiggy_CAC_Retention_Analytics.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

camp = pd.read_csv(f"{CLEAN}/campaign_performance_summary.csv")
chan = pd.read_csv(f"{CLEAN}/channel_performance_summary.csv")
rfm = pd.read_csv(f"{CLEAN}/customer_rfm_segments.csv")
member = pd.read_csv(f"{CLEAN}/membership_cac_comparison.csv")
repeat = pd.read_csv(f"{CLEAN}/membership_repeat_rate.csv")

SWIGGY_ORANGE = "FC8019"
HEADER_FILL = PatternFill("solid", fgColor=SWIGGY_ORANGE)
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="D9520A")
BASE_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

def style_header(ws, row, ncols):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

def write_df(ws, df, start_row=1, start_col=1):
    for j, col in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col+j, value=col)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, val in enumerate(row):
            cell = ws.cell(row=start_row+i, column=start_col+j, value=val)
            cell.font = BASE_FONT
            cell.border = BORDER
    style_header(ws, start_row, len(df.columns))
    for j, col in enumerate(df.columns):
        ws.column_dimensions[get_column_letter(start_col+j)].width = max(14, len(str(col))+4)
    return start_row + len(df) + 1

# ============ Sheet: Campaign_Data ============
ws1 = wb.active
ws1.title = "Campaign_Data"
write_df(ws1, camp)
ws1.freeze_panes = "A2"

# ============ Sheet: Channel_Data ============
ws2 = wb.create_sheet("Channel_Data")
write_df(ws2, chan)
ws2.freeze_panes = "A2"

# ============ Sheet: Membership_CAC (the core business question) ============
ws5 = wb.create_sheet("Membership_CAC")
last_m = write_df(ws5, member)
write_df(ws5, repeat, start_row=last_m+2)

# ============ Sheet: RFM_Segments ============
ws3 = wb.create_sheet("RFM_Segments")
write_df(ws3, rfm)
ws3.freeze_panes = "A2"
tier_col = get_column_letter(list(rfm.columns).index("tier")+1)
rule_champ = CellIsRule(operator="equal", formula=['"Champions"'], fill=PatternFill("solid", fgColor="C6EFCE"))
rule_lost = CellIsRule(operator="equal", formula=['"Lost"'], fill=PatternFill("solid", fgColor="FFC7CE"))
ws3.conditional_formatting.add(f"{tier_col}2:{tier_col}{len(rfm)+1}", rule_champ)
ws3.conditional_formatting.add(f"{tier_col}2:{tier_col}{len(rfm)+1}", rule_lost)
dv = DataValidation(type="list", formula1='"Champions,Loyal,Potential,At-Risk,Lost"', allow_blank=False)
ws3.add_data_validation(dv)
dv.add(f"{tier_col}2:{tier_col}{len(rfm)+1}")

# ============ Sheet: Dashboard ============
ws4 = wb.create_sheet("Dashboard", 0)
ws4.sheet_view.showGridLines = False
ws4["B2"] = "Swiggy Instamart — CAC & Retention Analytics (Case Study, Simulated Data)"
ws4["B2"].font = TITLE_FONT
ws4["B3"] = "Business question: does Swiggy One membership lower CAC and improve retention? All figures pull live via formulas."
ws4["B3"].font = Font(name="Arial", italic=True, size=9, color="666666")

n_camp, n_chan = len(camp), len(chan)

ws4["B5"] = "KPI"; ws4["C5"] = "Value"
style_header(ws4, 5, 2)

# Campaign_Data columns: A name,B sent,C opened,D clicked,E converted,F revenue,G spend,
#   H open_rate,I ctr,J conversion_rate,K roas,L cac
ws4["B6"], ws4["C6"] = "Total Spend (Rs)", f"=SUM(Campaign_Data!G2:G{n_camp+1})"
ws4["B7"], ws4["C7"] = "Total Revenue (Rs)", f"=SUM(Campaign_Data!F2:F{n_camp+1})"
ws4["B8"], ws4["C8"] = "Blended ROAS", "=IFERROR(C7/C6,0)"
ws4["B9"], ws4["C9"] = "Total Conversions", f"=SUM(Campaign_Data!E2:E{n_camp+1})"
ws4["B10"], ws4["C10"] = "Cheapest Channel (by CAC)", \
    f'=INDEX(Channel_Data!A2:A{n_chan+1},MATCH(MIN(Channel_Data!J2:J{n_chan+1}),Channel_Data!J2:J{n_chan+1},0))'
# Membership_CAC columns: A swiggy_one_member, B converted, C spend, D revenue, E interactions, F cac, G conversion_rate
ws4["B11"], ws4["C11"] = "Member CAC (Rs)", "=Membership_CAC!F3"
ws4["B12"], ws4["C12"] = "Non-Member CAC (Rs)", "=Membership_CAC!F2"
ws4["B13"], ws4["C13"] = "CAC Reduction from Membership", "=IFERROR((C12-C11)/C12,0)"

for r in range(6, 14):
    ws4.cell(row=r, column=2).font = BASE_FONT
    ws4.cell(row=r, column=3).font = Font(name="Arial", bold=True, size=11, color="D9520A")
    ws4.cell(row=r, column=2).border = BORDER
    ws4.cell(row=r, column=3).border = BORDER

ws4["C6"].number_format = '"Rs "#,##0'
ws4["C7"].number_format = '"Rs "#,##0'
ws4["C8"].number_format = '0.00"x"'
ws4["C9"].number_format = '#,##0'
ws4["C11"].number_format = '"Rs "#,##0'
ws4["C12"].number_format = '"Rs "#,##0'
ws4["C13"].number_format = '0.0%'

ws4.column_dimensions["A"].width = 3
ws4.column_dimensions["B"].width = 30
ws4.column_dimensions["C"].width = 20

# What-if lever: acquisition budget reallocation between two named PAID acquisition channels
# (excludes retention-only channels like the Swiggy One emailer / push, which don't compete
# for new-customer acquisition budget -- see docs/key_insights.md for why that comparison
# would be apples-to-oranges)
ws4["E5"] = "What-If: Shift Acquisition Budget"
ws4["E5"].font = Font(name="Arial", bold=True, size=11)
ws4["E6"] = "Budget to shift (Rs) ->"
ws4["F6"] = 50000
ws4["F6"].fill = PatternFill("solid", fgColor="FFFF00")
ws4["F6"].font = Font(name="Arial", bold=True)
ws4["E7"] = "From: Performance Ads CAC (Rs)"
ws4["F7"] = f'=INDEX(Channel_Data!J2:J{n_chan+1},MATCH("Performance Ads (Search/Display)",Channel_Data!A2:A{n_chan+1},0))'
ws4["E8"] = "To: Referral CAC (Rs)"
ws4["F8"] = f'=INDEX(Channel_Data!J2:J{n_chan+1},MATCH("Referral",Channel_Data!A2:A{n_chan+1},0))'
ws4["E9"] = "Extra conversions from shift"
ws4["F9"] = "=IFERROR(F6/F8 - F6/F7, 0)"
ws4["F9"].font = Font(name="Arial", bold=True, color="D9520A")
for r in [6,7,8,9]:
    ws4.cell(row=r, column=5).font = BASE_FONT
    ws4.cell(row=r, column=6).border = BORDER
ws4["F7"].number_format = '"Rs "#,##0'
ws4["F8"].number_format = '"Rs "#,##0'
ws4["F9"].number_format = '#,##0'
ws4.column_dimensions["E"].width = 26
ws4.column_dimensions["F"].width = 16

# Chart: CAC by channel
chart = BarChart()
chart.title = "CAC by Channel (lower is better)"
chart.y_axis.title = "CAC (Rs)"
chart.style = 10
data = Reference(ws2, min_col=10, min_row=1, max_row=n_chan+1)  # cac col (J)
cats = Reference(ws2, min_col=1, min_row=2, max_row=n_chan+1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.width, chart.height = 16, 8
ws4.add_chart(chart, "B15")

# Chart: Member vs Non-Member CAC
chart2 = BarChart()
chart2.title = "CAC: Swiggy One Members vs Non-Members"
chart2.y_axis.title = "CAC (Rs)"
chart2.style = 11
data2 = Reference(ws5, min_col=6, min_row=1, max_row=3)  # cac col in Membership_CAC
cats2 = Reference(ws5, min_col=1, min_row=2, max_row=3)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
chart2.width, chart2.height = 16, 8
ws4.add_chart(chart2, "B33")

# Data bars
cac_col_camp = get_column_letter(list(camp.columns).index("cac")+1)
ws1.conditional_formatting.add(
    f"{cac_col_camp}2:{cac_col_camp}{n_camp+1}",
    DataBarRule(start_type="min", end_type="max", color="FC8019")
)
cac_col_chan = get_column_letter(list(chan.columns).index("cac")+1)
ws2.conditional_formatting.add(
    f"{cac_col_chan}2:{cac_col_chan}{n_chan+1}",
    DataBarRule(start_type="min", end_type="max", color="FC8019")
)

wb.save(OUT)
print("Saved", OUT)
